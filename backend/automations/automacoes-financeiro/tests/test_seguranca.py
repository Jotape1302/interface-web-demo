"""Correções de segurança de 2026-10-06 (AUT-01..AUT-13): chave interna, baixa segura, NSU, segredos."""
import os
import stat
import threading
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
import requests
from fastapi.testclient import TestClient

from app import api as api_module
from app.api import app
from app.config import settings
from app.tiny_client import TinyClient

CHAVE = "k" * 40


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def com_chave(monkeypatch):
    monkeypatch.delenv("INTERNAL_AUTH_DISABLED", raising=False)
    monkeypatch.setattr(settings, "INTERNAL_AUTH_DISABLED", False)
    monkeypatch.setenv("AUTOMATIONS_INTERNAL_KEY", CHAVE)


# ---------------------------------------------------------------------------
# AUT-01: chave interna
# ---------------------------------------------------------------------------

def test_rota_protegida_sem_chave_401(client, com_chave):
    assert client.get("/api/v1/danfe/sefaz-status").status_code == 401
    assert client.get("/api/v1/danfe/sefaz-status", headers={"X-Internal-Key": "errada"}).status_code == 401


def test_rota_protegida_com_chave_passa(client, com_chave):
    with patch.object(api_module, "_build_sefaz_status_response", return_value={"ok": True}):
        res = client.get("/api/v1/conciliacao-ml/status", headers={"X-Internal-Key": CHAVE})
    assert res.status_code != 401 and res.status_code != 503


def test_post_de_baixa_sem_chave_nao_executa(client, com_chave):
    with patch.object(api_module, "_run_baixa") as run:
        res = client.post("/api/v1/reconciliacao/jobs/baixar", json={"itens": [{"empresa": "fabrilux", "conta_pagar_id": 1}]})
    assert res.status_code == 401
    run.assert_not_called()


def test_sem_chave_configurada_503(client, monkeypatch):
    monkeypatch.delenv("INTERNAL_AUTH_DISABLED", raising=False)
    monkeypatch.setattr(settings, "INTERNAL_AUTH_DISABLED", False)
    monkeypatch.delenv("AUTOMATIONS_INTERNAL_KEY", raising=False)
    monkeypatch.setattr(settings, "AUTOMATIONS_INTERNAL_KEY", None)
    res = client.get("/api/v1/danfe/sefaz-status", headers={"X-Internal-Key": CHAVE})
    assert res.status_code == 503
    assert "não configurada" in res.json()["detail"]


def test_health_publico_minimo_e_completo_com_chave(client, com_chave):
    publico = client.get("/health")
    assert publico.status_code == 200
    assert set(publico.json()) == {"status", "timestamp"}
    assert client.get("/api/v1/health").status_code == 200

    completo = client.get("/api/v1/health", headers={"X-Internal-Key": CHAVE}).json()
    assert completo["status"] == "ok"
    assert "nfe_provider" in completo and "tiny_configured" in completo


def test_health_tiny_exige_chave(client, com_chave):
    assert client.get("/health/tiny").status_code == 401


def test_cors_sem_origens_padrao_nao_libera_ip_nem_ddns(client):
    res = client.options(
        "/api/v1/danfe/sefaz-status",
        headers={"Origin": "http://192.168.50.20:3001", "Access-Control-Request-Method": "GET"},
    )
    assert "access-control-allow-origin" not in res.headers
    assert "access-control-allow-credentials" not in res.headers


# ---------------------------------------------------------------------------
# AUT-12: autor da auditoria vem do X-Usuario
# ---------------------------------------------------------------------------

def test_ml_baixa_usa_x_usuario_e_ignora_corpo(client):
    analise = SimpleNamespace(empresa_id="revesta", empresa_nome="Revesta")
    capturado = {}

    def baixar(analise_id, grupos, usuario="", progresso=None):
        capturado["usuario"] = usuario
        return {"status": "ok"}

    with patch.object(api_module, "_ml_analise_ou_404", return_value=analise), \
            patch.object(api_module, "ml_token_tiny", return_value="tok"), \
            patch.object(api_module.ml_rotina, "baixar", side_effect=baixar):
        job = client.post(
            "/api/v1/conciliacao-ml/jobs/baixar",
            json={"analise_id": "a1", "grupos": ["g1"], "usuario": "forjado"},
            headers={"X-Usuario": "Jo%C3%A3o"},
        ).json()
        for _ in range(100):
            if api_module._reconciliacao_jobs[job["job_id"]]["status"] != "executando":
                break
            threading.Event().wait(0.02)
    assert capturado["usuario"] == "João"


def test_usuario_do_corpo_vale_sem_header():
    from app.internal_auth import usuario_da_requisicao

    req = SimpleNamespace(headers={})
    assert usuario_da_requisicao(req, "maria") == "maria"
    assert usuario_da_requisicao(SimpleNamespace(headers={"X-Usuario": "ana"}), "maria") == "ana"


# ---------------------------------------------------------------------------
# AUT-02: baixa só do que a prévia deu como apto, com os valores do extrato
# ---------------------------------------------------------------------------

def _apto(conta_id=12345, inter_id="TX-1", valor="150.50", juros="0", desconto="0", valor_tiny=None, empresa="fabrilux"):
    api_module._registrar_aptos([SimpleNamespace(
        empresa=empresa, conta_pagar_id=conta_id, inter_transacao_id=inter_id,
        valor_pago=Decimal(valor), data_pagamento=date(2026, 9, 25),
        juros=Decimal(juros), multa=Decimal("0"), desconto=Decimal(desconto),
        valor_tiny=Decimal(valor_tiny or valor), historico_sugerido="Baixa conciliada Banco Inter",
    )])


def _conta(situacao="aberto", valor=150.50, saldo=None, valor_pago=0.0):
    return ({"id": 12345, "situacao": situacao, "valor": valor, "saldo": valor if saldo is None else saldo,
             "valorPago": valor_pago, "historico": "Ref. NF 1"}, "token-teste")


class _Tiny:
    def __init__(self, conta=None):
        self.baixar = MagicMock(return_value=True)
        self.obter = MagicMock(return_value=conta or _conta())
        self._patches = [
            patch.object(TinyClient, "baixar_conta_pagar_v3", self.baixar),
            patch.object(TinyClient, "obter_conta_pagar_v3", self.obter),
            patch.object(TinyClient, "resolve_conta_origem_inter", return_value=42),
            patch.object(settings, "TINY_API_TOKEN", "token-teste"),
            patch("app.api.TokenRepository"),
        ]

    def __enter__(self):
        for p in self._patches:
            m = p.start()
        m.return_value.is_configured.return_value = False
        return self

    def __exit__(self, *exc):
        for p in reversed(self._patches):
            p.stop()


def _item(**extra):
    return {"empresa": "fabrilux", "conta_pagar_id": 12345, "inter_transacao_id": "TX-1", **extra}


def test_baixa_recusa_item_fora_da_previa(client):
    with _Tiny() as tiny:
        data = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item()]}).json()
    assert data["total_sucessos"] == 0
    assert "Comparar" in data["resultados"][0]["mensagem"]
    tiny.baixar.assert_not_called()


def test_baixa_recusa_par_conta_debito_diferente_da_previa(client):
    _apto(inter_id="TX-1")
    with _Tiny() as tiny:
        data = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item(inter_transacao_id="TX-OUTRO")]}).json()
    assert data["resultados"][0]["status"] == "erro"
    tiny.baixar.assert_not_called()


def test_baixa_ignora_valores_do_cliente_e_usa_os_do_extrato(client):
    _apto(valor="152.00", juros="1.50", valor_tiny="150.50")
    forjado = _item(valor_pago=1.0, data_pagamento="2020-01-01", juros=0, desconto=999.0, conta_origem_id=666)
    with _Tiny() as tiny:
        data = client.post("/api/v1/reconciliacao/baixar", json={"itens": [forjado]}).json()
    assert data["total_sucessos"] == 1, data
    kw = tiny.baixar.call_args.kwargs
    assert kw["valor_pago"] == 152.00
    assert kw["data_pagamento"] == "2026-09-25"
    assert kw["juros"] == 1.50 and kw["desconto"] == 0.0
    assert kw["conta_origem_id"] == 42
    assert kw["conta_atual"]["situacao"] == "aberto"
    assert data["resultados"][0]["valor_pago"] == 152.00


def test_baixa_reler_conta_ja_paga_nao_baixa(client):
    _apto()
    with _Tiny(conta=_conta(situacao="pago", saldo=0, valor_pago=150.50)) as tiny:
        data = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item()]}).json()
    assert "já não está em aberto" in data["resultados"][0]["mensagem"]
    tiny.baixar.assert_not_called()


def test_baixa_reler_conta_com_valor_alterado_nao_baixa(client):
    _apto()
    with _Tiny(conta=_conta(valor=99.00)) as tiny:
        data = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item()]}).json()
    assert "mudou" in data["resultados"][0]["mensagem"]
    tiny.baixar.assert_not_called()


def test_baixa_maior_que_saldo_parcial_nao_baixa(client):
    _apto()
    with _Tiny(conta=_conta(situacao="parcial", saldo=50.0, valor_pago=100.50)) as tiny:
        data = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item()]}).json()
    assert "saldo em aberto" in data["resultados"][0]["mensagem"]
    tiny.baixar.assert_not_called()


def test_previa_registra_aptos_para_a_baixa(client):
    from app.inter_client import InterClient, InterTransacao

    extrato = [InterTransacao(
        empresa="fabrilux", transacao_id="TX-9", data_pagamento=date(2026, 9, 10),
        valor_pago=Decimal("500.00"), tipo_transacao="PIX", titulo="Pix enviado",
        descricao="ALUMIX", nome_favorecido="ALUMIX ALUMINIOS", cpf_cnpj_favorecido="11222333000181",
    )]

    def listar(token, situacao, **kw):
        if situacao == "aberto":
            return [{"id": 1, "numeroDocumento": "NF-1", "valor": 500.00, "dataVencimento": "2026-09-10",
                     "cliente": {"nome": "Alumix Aluminios", "cpfCnpj": "11.222.333/0001-81"}}]
        return []

    with patch.object(InterClient, "check_credentials_status", return_value={"status": "conectado", "mensagem": "ok"}), \
            patch.object(InterClient, "get_extrato", return_value=extrato), \
            patch.object(TinyClient, "list_contas_pagar_v3", side_effect=listar), \
            patch("app.api.TokenRepository") as repo, \
            patch.object(settings, "TINY_API_TOKEN", "token-teste"):
        repo.return_value.is_configured.return_value = False
        client.post("/api/v1/reconciliacao/preview", json={"empresa": "fabrilux", "data_inicio": "2026-09-01", "data_fim": "2026-09-28"})

    apto = api_module._apto_para("fabrilux", 1, "TX-9")
    assert apto is not None and apto.valor_pago == Decimal("500.00")


# ---------------------------------------------------------------------------
# AUT-03: baixa idempotente
# ---------------------------------------------------------------------------

def test_debito_ja_usado_nao_baixa_de_novo(client):
    _apto()
    with _Tiny() as tiny:
        primeira = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item()]}).json()
        segunda = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item()]}).json()
    assert primeira["total_sucessos"] == 1
    assert segunda["total_sucessos"] == 0
    assert "já usado" in segunda["resultados"][0]["mensagem"]
    assert tiny.baixar.call_count == 1


def test_item_repetido_no_mesmo_pedido(client):
    _apto()
    with _Tiny() as tiny:
        data = client.post("/api/v1/reconciliacao/baixar", json={"itens": [_item(), _item()]}).json()
    assert data["total_sucessos"] == 1 and data["total_falhas"] == 1
    assert tiny.baixar.call_count == 1


def test_job_de_baixa_simultaneo_da_mesma_conta_409(client):
    _apto()
    liberar = threading.Event()

    def baixar_devagar(**kw):
        liberar.wait(5)
        return True

    with _Tiny() as tiny:
        tiny.baixar.side_effect = baixar_devagar
        primeiro = client.post("/api/v1/reconciliacao/jobs/baixar", json={"itens": [_item()]})
        segundo = client.post("/api/v1/reconciliacao/jobs/baixar", json={"itens": [_item()]})
        liberar.set()
        for _ in range(200):
            if api_module._reconciliacao_jobs[primeiro.json()["job_id"]]["status"] != "executando":
                break
            threading.Event().wait(0.02)
    assert primeiro.status_code == 200
    assert segundo.status_code == 409
    assert api_module._baixas_em_andamento == set()


def _post_ok():
    return MagicMock(status_code=204)


def test_timeout_de_leitura_com_baixa_entrando_nao_repete_post():
    client = TinyClient()
    antes = {"situacao": "aberto", "valor": 100.0, "saldo": 100.0, "valorPago": 0.0, "historico": ""}
    depois = {"situacao": "pago", "valor": 100.0, "saldo": 0.0, "valorPago": 100.0}
    with patch("requests.post", side_effect=requests.ReadTimeout("read timeout")) as post, \
            patch.object(TinyClient, "obter_conta_pagar_v3", return_value=(depois, "t")), \
            patch("time.sleep"):
        assert client.baixar_conta_pagar_v3(token="t", conta_id=1, valor_pago=100.0, conta_atual=antes) is True
    assert post.call_count == 1


def test_timeout_de_leitura_sem_baixa_repete_depois_de_reler():
    client = TinyClient()
    antes = {"situacao": "aberto", "valor": 100.0, "saldo": 100.0, "valorPago": 0.0}
    with patch("requests.post", side_effect=[requests.ReadTimeout("x"), _post_ok()]) as post, \
            patch.object(TinyClient, "obter_conta_pagar_v3", return_value=(dict(antes), "t")) as obter, \
            patch("time.sleep"):
        assert client.baixar_conta_pagar_v3(token="t", conta_id=1, valor_pago=100.0, conta_atual=antes) is True
    assert post.call_count == 2
    assert obter.call_count == 1


def test_timeout_com_releitura_falhando_levanta_estado_incerto():
    client = TinyClient()
    with patch("requests.post", side_effect=requests.ReadTimeout("x")) as post, \
            patch.object(TinyClient, "obter_conta_pagar_v3", side_effect=RuntimeError("HTTP 503")), \
            patch("time.sleep"), \
            pytest.raises(RuntimeError, match="incerto"):
        client.baixar_conta_pagar_v3(token="t", conta_id=1, valor_pago=100.0, conta_atual={"situacao": "aberto"})
    assert post.call_count == 1


def test_http_5xx_e_ambiguo_e_rele_antes():
    client = TinyClient()
    depois = {"situacao": "pago", "valorPago": 100.0, "saldo": 0}
    with patch("requests.post", return_value=MagicMock(status_code=502)) as post, \
            patch.object(TinyClient, "obter_conta_pagar_v3", return_value=(depois, "t")), \
            patch("time.sleep"):
        assert client.baixar_conta_pagar_v3(
            token="t", conta_id=1, valor_pago=100.0, conta_atual={"situacao": "aberto", "valorPago": 0, "saldo": 100}
        ) is True
    assert post.call_count == 1


def test_falha_de_conexao_repete_sem_reler():
    client = TinyClient()
    with patch("requests.post", side_effect=[requests.ConnectTimeout("x"), _post_ok()]) as post, \
            patch.object(TinyClient, "obter_conta_pagar_v3") as obter, \
            patch("time.sleep"):
        assert client.baixar_conta_pagar_v3(token="t", conta_id=1, valor_pago=1.0, conta_atual={"situacao": "aberto"}) is True
    assert post.call_count == 2
    obter.assert_not_called()


def test_baixa_entrou_sem_foto_anterior():
    from app.tiny_client import baixa_entrou

    assert baixa_entrou(None, {"situacao": "pago"}) is True
    assert baixa_entrou(None, {"situacao": "aberto", "valorPago": 0}) is False
    assert baixa_entrou(None, {"situacao": "parcial", "valorPago": 30}) is None


# ---------------------------------------------------------------------------
# AUT-04: NSU validado
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("nsu", ["abc", "1</ultNSU><x>", "1234567890123456", "", " 12 3"])
def test_reset_nsu_invalido_422(client, nsu):
    with patch.object(api_module, "save_nsu_checkpoint") as salvar:
        res = client.post("/api/v1/danfe/sefaz-nsu/reset", json={"cnpj": "27184930000107", "set_nsu": nsu, "clear_cooldown": False})
    assert res.status_code == 422
    salvar.assert_not_called()


def test_reset_nsu_valido(client):
    with patch.object(api_module, "save_nsu_checkpoint") as salvar, \
            patch.object(api_module, "load_nsu_state", return_value={}), \
            patch.object(api_module, "load_cooldown_state", return_value={}):
        res = client.post("/api/v1/danfe/sefaz-nsu/reset", json={"cnpj": "27.184.930/0001-07", "set_nsu": "000000000012345", "clear_cooldown": False})
    assert res.status_code == 200
    salvar.assert_called_once_with("27184930000107", "000000000012345", settings)


def test_reset_cnpj_invalido_422(client):
    assert client.post("/api/v1/danfe/sefaz-nsu/reset", json={"cnpj": "../../x", "clear_cooldown": True}).status_code == 422


def test_envelope_soap_recusa_nsu_com_xml():
    from app.nfe_client import SefazA1Provider

    provider = SefazA1Provider(cnpj="27184930000107")
    with pytest.raises(ValueError):
        provider._build_dist_nsu_envelope("1</ultNSU><evil/>")
    assert "<ultNSU>000000000000042</ultNSU>" in provider._build_dist_nsu_envelope("42")


# ---------------------------------------------------------------------------
# AUT-05 / AUT-09 / AUT-11 / AUT-13
# ---------------------------------------------------------------------------

def test_cache_oauth_do_tiny_gravado_0600(tmp_path):
    from app.tiny_client import TinyOAuthManager

    cache = tmp_path / "tiny_token_cache.json"
    mgr = TinyOAuthManager(client_id="id", client_secret="s", cache_file=cache)
    mgr._save_cache({"access_token": "a", "refresh_token": "r", "expires_in": 100})
    assert stat.S_IMODE(os.stat(cache).st_mode) == 0o600


def test_estado_comum_continua_0644(tmp_path):
    from app.json_store import write_json_atomic

    destino = tmp_path / "estado.json"
    write_json_atomic(destino, {"a": 1})
    assert stat.S_IMODE(os.stat(destino).st_mode) == 0o644


def test_sem_token_e_sem_mock_mode_nao_simula_lancamento(monkeypatch):
    from app.categorization import GroupedLancamento
    from app.nfe_client import MockNFeProvider
    from app.tiny_client import TinyAPIError

    monkeypatch.setattr(settings, "MOCK_MODE", False)
    nfe = MockNFeProvider().get_nfe_data("35230912345678000123550010000001231000000123")
    grupo = MagicMock(spec=GroupedLancamento)
    grupo.valor_total = Decimal("10")
    grupo.categoria_nome = "Teste"
    client = TinyClient(token="mock")
    with patch.object(client, "resolve_token_for_nfe", return_value=("mock", None)), \
            patch.object(client, "_resolve_and_validate_marcadores", return_value=[]), \
            patch.object(client, "_resolve_doc_base", return_value="NF-1"), \
            pytest.raises(TinyAPIError, match="NÃO foi feito"):
        client.lancar_grupo_categorizado(nfe, grupo, 1, 1)


def test_upload_maior_que_o_limite_e_recusado(client, monkeypatch):
    monkeypatch.setattr(settings, "UPLOAD_MAX_FILE_MB", 1)
    grande = b"x" * (1024 * 1024 + 10)
    with patch.object(api_module, "DANFEScanner") as scanner, patch.object(api_module, "get_nfe_provider"):
        res = client.post("/api/v1/danfe/process-batch", files=[("files", ("grande.pdf", grande, "application/pdf"))])
    data = res.json()
    assert data["failed_count"] == 1
    assert "limite" in data["items"][0]["error_message"]
    scanner.return_value.scan_file.assert_not_called()


def test_pdf_com_paginas_demais_e_recusado(client, monkeypatch):
    monkeypatch.setattr(settings, "UPLOAD_MAX_PDF_PAGES", 50)
    with patch.object(api_module, "DANFEScanner") as scanner, patch.object(api_module, "get_nfe_provider"), \
            patch.object(api_module, "_paginas_pdf", return_value=51):
        res = client.post("/api/v1/danfe/process-batch", files=[("files", ("n.pdf", b"%PDF-1.4", "application/pdf"))])
    assert "51 páginas" in res.json()["items"][0]["error_message"]
    scanner.return_value.scan_file.assert_not_called()


def test_lote_com_arquivos_demais_413(client, monkeypatch):
    monkeypatch.setattr(settings, "UPLOAD_MAX_FILES", 2)
    arquivos = [("files", (f"{i}.xml", b"<a/>", "text/xml")) for i in range(3)]
    assert client.post("/api/v1/danfe/process-batch", files=arquivos).status_code == 413


def test_pem_da_a1_em_pasta_privada_e_apagado(tmp_path):
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from cryptography.x509.oid import NameOID
    from datetime import datetime, timedelta, timezone

    from app.nfe_client import SefazA1Provider

    chave = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    nome = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "teste")])
    agora = datetime.now(timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(nome).issuer_name(nome).public_key(chave.public_key())
            .serial_number(1).not_valid_before(agora).not_valid_after(agora + timedelta(days=1))
            .sign(chave, hashes.SHA256()))
    pfx = tmp_path / "a1.pfx"
    pfx.write_bytes(pkcs12.serialize_key_and_certificates(b"t", chave, cert, None, serialization.BestAvailableEncryption(b"s")))

    provider = SefazA1Provider(cert_path=pfx, cert_password="s", cnpj="27184930000107")
    cert_file, key_file = provider._prepare_certificates()
    pasta = os.path.dirname(key_file)
    assert stat.S_IMODE(os.stat(pasta).st_mode) == 0o700
    assert stat.S_IMODE(os.stat(key_file).st_mode) == 0o600
    provider._cleanup_temp_certs()
    assert not os.path.exists(pasta)


def test_pem_e_apagado_mesmo_quando_a_consulta_falha(tmp_path):
    from app.nfe_client import NFeProviderError, SefazA1Provider

    provider = SefazA1Provider(cert_path=tmp_path / "x.pfx", cnpj="27184930000107")
    pastas = []

    def preparar():
        pastas.append(api_module.tempfile.mkdtemp())
        provider._temp_dir = pastas[-1]
        return ("c", "k")

    with patch.object(provider, "_prepare_certificates", side_effect=preparar), \
            patch("app.nfe_client.requests.post", side_effect=requests.ConnectionError("x")), \
            pytest.raises(NFeProviderError):
        provider.consultar_distribuicao_nsu("1")
    assert pastas and not os.path.exists(pastas[0])


# ---------------------------------------------------------------------------
# Reauditoria de 2026-10-07 (SEC-03, SEC-04, SEC-05)
# ---------------------------------------------------------------------------

def test_sec03_planilha_ml_nao_grava_texto_como_formula():
    import io

    from openpyxl import load_workbook

    from app import ml_reconciliation as ml

    conteudo = ml._planilha(
        ["Nome do cliente", "Valor"],
        [['=HYPERLINK("http://exemplo.invalid";"Clique")', Decimal("10.50")], ["+5511999", Decimal("-1")]],
    )
    ws = load_workbook(io.BytesIO(conteudo)).active
    assert ws["A2"].data_type == "s" and ws["A2"].value == '=HYPERLINK("http://exemplo.invalid";"Clique")'
    assert ws["A3"].value == "+5511999"
    assert ws["B2"].value == 10.5 and ws["B3"].value == -1


def test_sec04_fila_de_tarefas_separada_entre_inter_e_ml(client, monkeypatch):
    monkeypatch.setattr(settings, "INTERNAL_AUTH_DISABLED", True)
    inter = api_module._iniciar_job("baixar", lambda progresso: {"ok": True}, sincrono=True)["job_id"]
    ml = api_module._iniciar_job("ml_baixar", lambda progresso: {"ok": True}, sincrono=True)["job_id"]
    assert client.get(f"/api/v1/reconciliacao/jobs/{inter}").status_code == 200
    assert client.get(f"/api/v1/conciliacao-ml/jobs/{ml}").status_code == 200
    assert client.get(f"/api/v1/reconciliacao/jobs/{ml}").status_code == 404
    assert client.get(f"/api/v1/conciliacao-ml/jobs/{inter}").status_code == 404


def test_sec05_documentacao_da_api_desligada(client, com_chave):
    for rota in ("/docs", "/redoc", "/openapi.json"):
        assert client.get(rota).status_code in (401, 404), rota
        assert client.get(rota, headers={"X-Internal-Key": CHAVE}).status_code == 404, rota
