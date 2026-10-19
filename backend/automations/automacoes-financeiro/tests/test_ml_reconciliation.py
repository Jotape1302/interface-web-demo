import io
import time
from datetime import date
from decimal import Decimal
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import ml_reconciliation as ml
from app import ml_reconciliation_rotina as rot
from app.config import Settings
from app.ml_apis import (
    APIMercadoError,
    ClienteMercadoLivre,
    ClienteMercadoPago,
    SemCredencial,
    ler_csv_liberacoes,
    limites_periodo,
)

NORVIA = ml.EMPRESAS_ML["norvia_campinas"]


def _lib(**campos):
    base = {
        "DATE": "2025-09-01T10:00:00.000-03:00",
        "RECORD_TYPE": "release",
        "DESCRIPTION": "payment",
        "NET_CREDIT_AMOUNT": "0",
        "NET_DEBIT_AMOUNT": "0",
        "GROSS_AMOUNT": "0",
        "ORDER_ID": "",
        "PACK_ID": "",
    }
    base.update({k: str(v) for k, v in campos.items()})
    return base


# Liberações do Mercado Pago (relatório de Liberações, já lido do CSV).
LIBERACOES = [
    _lib(RECORD_TYPE="initial_available_balance", NET_CREDIT_AMOUNT="1000"),
    # Venda da Ana (pacote 3000...1, pedido 2000...1): bruto 110, tarifa 12,50, frete 7,50 -> líquido 90
    _lib(SOURCE_ID="5001", NET_CREDIT_AMOUNT="90.00", GROSS_AMOUNT="110.00", MP_FEE_AMOUNT="-12.50",
         SHIPPING_FEE_AMOUNT="-7.50", ORDER_ID="2000000000000001"),
    # Venda do Bruno, localizada só pelo pedido do relatório (pagamento não consta no pedido do ML)
    _lib(SOURCE_ID="5003", NET_CREDIT_AMOUNT="80.00", GROSS_AMOUNT="80.00", ORDER_ID="2000000000000002"),
    # Liberação que não fecha: bruto 50, tarifa 5, líquido 44 -> diferença de -1
    _lib(SOURCE_ID="5004", NET_CREDIT_AMOUNT="44.00", GROSS_AMOUNT="50.00", MP_FEE_AMOUNT="-5.00",
         ORDER_ID="2000000000000003"),
    # Dinheiro retido e liberado no mesmo pagamento: anula
    _lib(SOURCE_ID="5005", DESCRIPTION="reserve_for_dispute", NET_DEBIT_AMOUNT="30.00"),
    _lib(SOURCE_ID="5005", DESCRIPTION="reserve_for_dispute", NET_CREDIT_AMOUNT="30.00"),
    # Publicidade sem par: verificar
    _lib(SOURCE_ID="5006", DESCRIPTION="payment", NET_DEBIT_AMOUNT="20.00"),
    # Saque: fora dos lançamentos
    _lib(SOURCE_ID="5007", DESCRIPTION="withdrawal", NET_DEBIT_AMOUNT="500.00"),
    # Tipo desconhecido
    _lib(SOURCE_ID="5008", DESCRIPTION="novidade_mp", NET_CREDIT_AMOUNT="3.00"),
    # Liberação sem venda no ML nem pedido no relatório
    _lib(SOURCE_ID="5999", NET_CREDIT_AMOUNT="30.00", GROSS_AMOUNT="30.00", DATE="2025-09-02T01:30:00Z"),
]
# Linha repetida no relatório: vira pendência de conferência (não some, não baixa de novo)
LIBERACOES.append(dict(LIBERACOES[-1]))

PEDIDOS_ML = [
    {"id": 2000000000000001, "pack_id": 3000000000000001, "status": "paid", "payments": [{"id": 5001}]},
    {"id": 2000000000000002, "pack_id": None, "status": "paid", "payments": [{"id": 9999}]},
    {"id": 2000000000000003, "pack_id": None, "status": "paid", "payments": [{"id": 5004}]},
]

PEDIDOS_TINY = [
    {"ecommerce": {"numeroPedidoEcommerce": "3000000000000001"}, "cliente": {"nome": "Ana Souza", "cpfCnpj": "111.111.111-11"}},
    {"ecommerce": {"numeroPedidoEcommerce": "2000000000000002"}, "cliente": {"nome": "Bruno Lima", "cpfCnpj": "222.222.222-22"}},
    {"ecommerce": {"numeroPedidoEcommerce": "2000000000000003"}, "cliente": {"nome": "Carla Dias", "cpfCnpj": ""}},
]


def _analisar(listar=None, baixas=None):
    return ml.analisar("a1", NORVIA, LIBERACOES, PEDIDOS_ML, PEDIDOS_TINY, listar,
                       (date(2025, 9, 1), date(2025, 9, 2)), baixas_feitas=baixas)


# ─── regras ───────────────────────────────────────────────────────────────────


def test_data_sp_converte_utc_para_sao_paulo():
    assert ml.data_sp("2025-09-01T01:50:29Z") == date(2025, 8, 31)
    assert ml.data_sp("2025-09-01T10:00:00.000-03:00") == date(2025, 9, 1)
    assert ml.data_sp("15/09/2025") == date(2025, 9, 15)
    assert ml.limpar_id(2000013231387746.0) == "2000013231387746"
    assert ml.limpar_id("1.23868190197E+11") == "123868190197"


def test_expandir_liberacoes_decompoe_venda_descarta_saldo_e_confere_duplicata():
    movs = ml.expandir_liberacoes(LIBERACOES)
    venda = [(m.tipo_operacao, m.valor) for m in movs if m.id_pagamento == "5001"]
    assert venda == [("Recebimento", Decimal("110.00")), ("Tarifa de venda", Decimal("-12.50")), ("Tarifa de envio", Decimal("-7.50"))]
    assert [(m.tipo_operacao, m.valor) for m in movs if m.id_pagamento == "5004"][-1] == (ml.TIPO_DIFERENCA, Decimal("-1.00"))
    assert [m.tipo_operacao for m in movs if m.id_pagamento == "5999"] == [ml.TIPO_RECEBIMENTO, ml.TIPO_LINHA_REPETIDA]
    assert {m.tipo_operacao for m in movs if m.id_pagamento == "5005"} == {"Dinheiro retido", "Liberação de dinheiro retido"}
    assert [m.tipo_operacao for m in movs if m.id_pagamento == "5007"] == [ml.TIPO_TRANSFERENCIA]
    assert next(m for m in movs if m.id_pagamento == "5999").data == "01/09/2025"  # 01:30 UTC = 22:30 de 01/09 em SP


def test_cruzar_vendas_pelo_pagamento_pelo_pedido_e_sem_venda():
    movs = ml.expandir_liberacoes(LIBERACOES)
    comp = {c.id_pagamento: c for c in ml.cruzar_vendas(movs, PEDIDOS_ML, PEDIDOS_TINY)}
    assert comp["5001"].casou_por == ml.CASOU_PAGAMENTO
    assert (comp["5001"].id_pacote, comp["5001"].nome_cliente) == ("3000000000000001", "Ana Souza")
    assert comp["5003"].casou_por == ml.CASOU_PEDIDO and "só pelo pedido" in comp["5003"].observacao
    assert comp["5999"].casou_por == ml.CASOU_NENHUM
    assert comp["5001"].valor_liberado == Decimal("90.00")
    ana = next(m for m in movs if m.id_pagamento == "5001")
    assert (ana.id_venda, ana.cpf_cnpj) == ("3000000000000001", "111.111.111-11")


def test_divergencia_entre_pedido_do_relatorio_e_do_pagamento():
    libs = [_lib(SOURCE_ID="5001", NET_CREDIT_AMOUNT="10", GROSS_AMOUNT="10", ORDER_ID="2000000000000777")]
    movs = ml.expandir_liberacoes(libs)
    (c,) = ml.cruzar_vendas(movs, PEDIDOS_ML, PEDIDOS_TINY)
    assert c.casou_por == ml.CASOU_PAGAMENTO and "diferente do pedido do pagamento" in c.observacao


def test_selecionar_conta_prefere_parcela_aberta_mais_baixa():
    contas = [
        {"id": 1, "historico": "OC 123 parcela 1/2", "situacao": "pago"},
        {"id": 2, "historico": "OC 123 parcela 2/2", "situacao": "aberto"},
        {"id": 3, "historico": "OC 1234 parcela 1/1", "situacao": "aberto"},
    ]
    assert ml.selecionar_conta(contas, "123")["id"] == 2
    assert ml.selecionar_conta(contas[:1], "123")["id"] == 1
    assert ml.selecionar_conta(contas, "999") is None


def test_avaliar_conta_valor_acima_do_saldo_nao_e_apto():
    g = ml.GrupoRecebimento(id="G", id_venda="1", indices=[0], chaves=["k"], valor_total=Decimal("110"), data="", nome_cliente="x")
    ml.aplicar_conta_ao_grupo(g, {"id": 9, "situacao": "aberto", "valor": 100, "saldo": 100})
    ml.avaliar_conta(g)
    assert g.situacao == ml.SIT_VALOR_DIVERGENTE and not g.apto_baixa
    ml.aplicar_conta_ao_grupo(g, {"id": 9, "situacao": "parcial", "valor": 200, "saldo": 150})
    ml.avaliar_conta(g)
    assert g.apto_baixa and "Baixa parcial" in g.mensagem


def test_analise_completa_sem_tiny():
    a = _analisar()
    res = ml.resumo(a)
    assert res["total_grupos"] == 4  # Ana, Bruno, Carla, sem venda
    assert res["grupos_por_situacao"][ml.SIT_SEM_CLIENTE] == 1
    assert res["grupos_por_situacao"][ml.SIT_TINY_NAO_CONSULTADO] == 3
    assert res["total_anuladas"] == 2
    motivos = {p.tipo_operacao: p.motivo for p in a.lancamentos.pendencias}
    assert "sem par" in motivos["Pagamento"]
    assert "não fecha" in motivos[ml.TIPO_DIFERENCA]
    assert "sem categoria" in motivos["Outro movimento do Mercado Pago (novidade_mp)"]
    assert "CPF/CNPJ" in motivos["Tarifa de venda"]  # Carla sem CPF no Tiny
    lanc = {(l.tipo_operacao, l.contato): l for l in a.lancamentos.lancamentos}
    tarifa = lanc[("Tarifa de venda", "Ana Souza")]
    assert (tarifa.tipo, tarifa.valor, tarifa.cnpj, tarifa.numero_documento) == ("D", Decimal("12.50"), "111.111.111-11", "5001")
    assert tarifa.historico.startswith("Ref as tarifas sobre a venda 3000000000000001, operação 5001")
    assert lanc[("Tarifa de envio", "Ana Souza")].categoria == "Fretes sobre Vendas"


def test_analise_com_tiny_e_baixas_ja_feitas():
    contas = {
        "Ana Souza": [{"id": 10, "historico": "OC 3000000000000001", "situacao": "aberto", "valor": 110, "saldo": 110}],
        "Bruno Lima": [{"id": 20, "historico": "OC 2000000000000002", "situacao": "pago", "valor": 80, "saldo": 0}],
        "Carla Dias": [],
    }
    a = _analisar(listar=lambda nome: contas[nome])
    sit = {g.id_venda: g.situacao for g in a.grupos}
    assert sit == {
        "3000000000000001": ml.SIT_APTO_BAIXA,
        "2000000000000002": "STATUS: PAGO",
        "2000000000000003": ml.SIT_NAO_LOCALIZADO,
        "": ml.SIT_SEM_CLIENTE,
    }
    chave_ana = next(g for g in a.grupos if g.id_venda == "3000000000000001").chaves[0]
    b = _analisar(listar=lambda nome: contas[nome], baixas={chave_ana: "2025-09-03"})
    ana = next(g for g in b.grupos if g.id_venda == "3000000000000001")
    assert ana.situacao == ml.SIT_CONCILIADO and ana.baixado_em == "2025-09-03" and not ana.apto_baixa


def test_persistencia_ida_e_volta_preserva_centavos():
    a = _analisar()
    b = ml.analise_de_dict(ml.analise_para_dict(a))
    assert ml.analise_json(a) == ml.analise_json(b)
    assert isinstance(b.movimentos[0].valor, Decimal)


def test_planilhas_para_download():
    a = _analisar()
    for tipo in ml.TIPOS_ARQUIVO:
        nome, conteudo = ml.gerar_arquivo(a, tipo)
        assert nome.endswith(".xlsx") and "NORVIA CAMPINAS" in nome
        ws = load_workbook(io.BytesIO(conteudo)).active
        assert ws.max_row >= 2
    ws = load_workbook(io.BytesIO(ml.gerar_arquivo(a, "lancamentos")[1])).active
    assert [c.value for c in ws[1]][:5] == ["Data", "Categoria", "Histórico", "Tipo", "Valor"]


# ─── APIs do Mercado Pago / Mercado Livre ─────────────────────────────────────


class _Resp:
    def __init__(self, status=200, json_data=None, content=b""):
        self.status_code = status
        self._json = json_data
        self.content = content if content else (b"x" if json_data is not None else b"")
        self.text = str(json_data)

    def json(self):
        return self._json


class _HTTP:
    def __init__(self, rotas):
        self.rotas = rotas
        self.chamadas = []

    def request(self, metodo, url, **kw):
        self.chamadas.append((metodo, url, kw.get("params"), kw.get("json")))
        for (m, trecho), resposta in self.rotas.items():
            if m == metodo and url.endswith(trecho):
                return resposta(kw) if callable(resposta) else resposta
        raise AssertionError(f"rota inesperada {metodo} {url}")


CSV = (
    "DATE;SOURCE_ID;DESCRIPTION;NET_CREDIT_AMOUNT;NET_DEBIT_AMOUNT;GROSS_AMOUNT;RECORD_TYPE\n"
    "2025-09-01T10:00:00.000-03:00;5001;payment;90.00;0;110.00;release\n"
).encode()


def test_ler_csv_liberacoes_exige_colunas():
    assert ler_csv_liberacoes(CSV)[0]["SOURCE_ID"] == "5001"
    with pytest.raises(APIMercadoError, match="sem as colunas"):
        ler_csv_liberacoes(b"FECHA;ID\n1;2\n")


def test_limites_periodo_em_utc():
    assert limites_periodo(date(2025, 9, 1), date(2025, 9, 1)) == ("2025-09-01T03:00:00Z", "2025-09-02T02:59:59Z")


def test_cliente_mp_configura_pede_espera_e_baixa_relatorio():
    listagens = [[], [{"id": 77, "file_name": "rel.csv", "status": "processed"}]]
    http = _HTTP({
        ("GET", "/release_report/config"): _Resp(404, {"message": "not found"}),
        ("POST", "/release_report/config"): _Resp(201, {}),
        ("POST", "/v1/account/release_report"): _Resp(202, {"id": 77}),
        ("GET", "/release_report/list"): lambda kw: _Resp(200, listagens.pop(0)),
        ("GET", "/release_report/rel.csv"): _Resp(200, content=CSV),
    })
    esperas = []
    cli = ClienteMercadoPago("norvia_campinas", lambda: "tok", http=http, cfg=Settings(), dormir=esperas.append)
    linhas = cli.liberacoes(date(2025, 9, 1), date(2025, 9, 1))
    assert linhas[0]["NET_CREDIT_AMOUNT"] == "90.00" and len(esperas) == 1
    pedido = next(c for c in http.chamadas if c[0] == "POST" and c[1].endswith("/v1/account/release_report"))
    assert pedido[3] == {"begin_date": "2025-09-01T03:00:00Z", "end_date": "2025-09-02T02:59:59Z"}
    config = next(c for c in http.chamadas if c[0] == "POST" and c[1].endswith("/config"))
    assert {"key": "SOURCE_ID"} in config[3]["columns"] and config[3]["display_timezone"] == "GMT-03"


def test_cliente_mp_sem_token_e_401():
    with pytest.raises(SemCredencial):
        ClienteMercadoPago("revesta", lambda: None, http=_HTTP({}), cfg=Settings()).garantir_configuracao()
    http = _HTTP({("GET", "/release_report/config"): _Resp(401, {})})
    with pytest.raises(SemCredencial, match="401"):
        ClienteMercadoPago("revesta", lambda: "tok", http=http, cfg=Settings()).garantir_configuracao()


def test_cliente_ml_pagina_pedidos_em_janelas():
    paginas = {0: {"results": [{"id": 1}, {"id": 2}], "paging": {"total": 3}}, 2: {"results": [{"id": 3}], "paging": {"total": 3}}}

    def busca(kw):
        return _Resp(200, paginas.get(kw["params"]["offset"], {"results": [], "paging": {"total": 0}}))

    http = _HTTP({("GET", "/users/me"): _Resp(200, {"id": 42}), ("GET", "/orders/search"): busca})
    cli = ClienteMercadoLivre("norvia", lambda: "tok", http=http, cfg=Settings())
    cli.LIMITE_PAGINA = 2
    pedidos = cli.pedidos(date(2025, 9, 1), date(2025, 9, 3))
    assert [p["id"] for p in pedidos] == [1, 2, 3]
    primeira = next(c for c in http.chamadas if c[1].endswith("/orders/search"))
    assert primeira[2]["seller"] == 42 and primeira[2]["order.date_created.from"] == "2025-09-01T00:00:00.000-03:00"


# ─── rotina (execução, histórico, avisos, baixa) ──────────────────────────────


class _MP:
    def liberacoes(self, de, ate, progresso=None):
        return LIBERACOES


class _ML:
    def pedidos(self, de, ate, progresso=None):
        return PEDIDOS_ML[:2]

    def pedido(self, pid):
        return next((p for p in PEDIDOS_ML if str(p["id"]) == pid), None)


class _MLSemToken(_ML):
    def pedidos(self, de, ate, progresso=None):
        raise SemCredencial("Mercado Livre: token expirado ou inválido (HTTP 401).")


class _Tiny:
    def __init__(self):
        self.baixas = []
        self.contas = {
            "Ana Souza": [{"id": 10, "historico": "OC 3000000000000001", "situacao": "aberto", "valor": 110, "saldo": 110}],
            "Bruno Lima": [{"id": 20, "historico": "OC 2000000000000002", "situacao": "aberto", "valor": 80, "saldo": 80}],
        }
        self.situacao_atual = {10: "aberto", 20: "pago"}

    def list_pedidos_v3(self, token, data_inicial=None, data_final=None, numero_pedido_ecommerce=None, max_pages=300):
        if numero_pedido_ecommerce:
            return [p for p in PEDIDOS_TINY if p["ecommerce"]["numeroPedidoEcommerce"] == numero_pedido_ecommerce]
        return PEDIDOS_TINY[:2]

    def list_contas_receber_v3(self, token, nome_cliente=None, **kw):
        return self.contas.get(nome_cliente, [])

    def obter_conta_receber_v3(self, token, cid):
        sit = self.situacao_atual[cid]
        return {"situacao": sit, "valor": 110 if cid == 10 else 80, "saldo": 0 if sit == "pago" else 110}

    def baixar_conta_receber_v3(self, token, cid, **kw):
        self.baixas.append((cid, kw))
        self.situacao_atual[cid] = "pago"

    def obter_contas_financeiras_v3(self, token):
        return [{"id": 555, "descricao": "Mercado Livre - REVESTA"}, {"id": 556, "descricao": "Inter"}]

    def list_categorias_receita_despesa_v3(self, token):
        return [{"id": 882461346, "descricao": "Vendas Mercado Livre"}] if token == "tok-norvia_campinas" else [
            {"id": 900, "descricao": "Vendas  mercado livre"}, {"id": 901, "descricao": "Outras"}]


def _rotina(tmp_path, ml_cliente=_ML, tiny=None):
    cfg = Settings(DATA_DIR=tmp_path)
    return rot.RotinaConciliacaoML(
        estado=rot.EstadoML(cfg=cfg),
        cfg=cfg,
        tiny=tiny or _Tiny(),
        fabrica_mp=lambda e: _MP(),
        fabrica_ml=lambda e: ml_cliente(),
        obter_token_tiny=lambda emp: f"tok-{emp}",
    )


def test_rotina_executa_registra_log_aviso_e_analise(tmp_path):
    r = _rotina(tmp_path)
    ex = r.executar("agendada", ["norvia_campinas"], date(2025, 9, 1), date(2025, 9, 2))
    assert ex["status"] == "ok" and ex["tipo"] == "analise"
    res = ex["empresas"]["norvia_campinas"]["resumo"]
    assert res["aptos_baixa"] == 2
    mensagens = " ".join(l["mensagem"] for l in ex["log"])
    assert "1 venda(s) anterior(es)" in mensagens  # pedido 2000...3 buscado pelo relatório
    assert "pedido(s) avulso(s)" in mensagens  # Carla buscada no Tiny pelo número
    dados = r.estado.ler()
    assert dados["avisos"][-1]["tipo"] == "conciliacao_ml_pronta"
    assert "2 apto(s) para baixa (R$ 190,00)" in dados["avisos"][-1]["mensagem"]
    analise = r.estado.analise(ex["empresas"]["norvia_campinas"]["analise_id"])
    assert analise is not None and analise.origem == "agendada"
    assert (tmp_path / "conciliacao_ml" / "estado.json").exists()

    # Sobrevive a um "restart": outro EstadoML lê do disco.
    outro = rot.EstadoML(cfg=r.cfg)
    assert outro.analise(analise.id).empresa_id == "norvia_campinas"
    st = rot.RotinaConciliacaoML(estado=outro, cfg=r.cfg, tiny=_Tiny(), obter_token_tiny=lambda e: None).status()
    assert st["ultima_execucao"]["linhas_log"] > 0 and "log" not in st["ultima_execucao"]
    assert st["empresas"][0]["ultima_analise"]["analise_id"] == analise.id


def test_rotina_sem_token_ml_segue_pelo_pedido_do_relatorio(tmp_path):
    r = _rotina(tmp_path, ml_cliente=_MLSemToken)
    ex = r.executar("manual", ["norvia_campinas"], date(2025, 9, 1), date(2025, 9, 1))
    assert ex["status"] == "ok"
    assert any(l["nivel"] == "aviso" and "Vendas do Mercado Livre não lidas" in l["mensagem"] for l in ex["log"])
    analise = r.estado.analise(ex["empresas"]["norvia_campinas"]["analise_id"])
    assert {c.casou_por for c in analise.comparativo} <= {ml.CASOU_PEDIDO, ml.CASOU_NENHUM}


def test_rotina_falha_numa_empresa_nao_para_a_outra(tmp_path):
    class _MPQuebrado:
        def liberacoes(self, *a, **k):
            raise APIMercadoError("Mercado Pago: falha ao pedir o relatório: HTTP 500")

    r = _rotina(tmp_path)
    r.fabrica_mp = lambda e: _MPQuebrado() if e.id == "revesta" else _MP()
    ex = r.executar("agendada", None, date(2025, 9, 1), date(2025, 9, 1))
    assert ex["status"] == "parcial"
    assert ex["empresas"]["revesta"]["status"] == "erro"
    assert r.estado.ler()["avisos"][0]["tipo"] == "conciliacao_ml_erro"


def test_rotina_nao_roda_duas_vezes_ao_mesmo_tempo(tmp_path):
    r = _rotina(tmp_path)
    r._executando.acquire()
    with pytest.raises(rot.EmExecucao):
        r.executar("manual", ["norvia_campinas"], date(2025, 9, 1), date(2025, 9, 1))


def test_validar_periodo():
    rot.validar_periodo(date(2025, 9, 1), date(2025, 9, 30), hoje=date(2025, 10, 1))
    for de, ate in ((date(2025, 9, 2), date(2025, 9, 1)), (date(2025, 9, 1), date(2025, 10, 2)), (date(2025, 7, 1), date(2025, 9, 30))):
        with pytest.raises(ValueError):
            rot.validar_periodo(de, ate, hoje=date(2025, 10, 1))
    assert rot.periodo_padrao(date(2025, 10, 5)) == (date(2025, 8, 6), date(2025, 10, 4))


def test_baixa_reconfere_registra_e_nao_repete(tmp_path):
    tiny = _Tiny()
    r = _rotina(tmp_path, tiny=tiny)
    ex = r.executar("manual", ["norvia_campinas"], date(2025, 9, 1), date(2025, 9, 2))
    analise_id = ex["empresas"]["norvia_campinas"]["analise_id"]
    aptos = [g.id for g in r.estado.analise(analise_id).grupos if g.apto_baixa]
    assert len(aptos) == 2

    res = r.baixar(analise_id, aptos, usuario="Fulano")
    # Bruno foi pago por outra pessoa entre a análise e a baixa: pulado.
    assert (res["total_sucessos"], res["total_ignorados"], res["total_falhas"]) == (1, 1, 0)
    assert tiny.baixas == [(10, {"valor_pago": 110.0, "data_pagamento": "01/09/2025",
                                 "conta_destino_id": 702641458, "categoria_id": 882461346})]
    salva = r.estado.analise(analise_id)
    assert next(g for g in salva.grupos if g.conta_id == 10).situacao == ml.SIT_CONCILIADO
    historico = r.estado.ler()["execucoes"][-1]
    assert historico["tipo"] == "baixa" and historico["usuario"] == "Fulano" and historico["status"] == "ok"

    # Reprocessar o mesmo período com a conta ainda "parcial" no Tiny não reabre a baixa.
    tiny.situacao_atual[10] = "parcial"
    tiny.contas["Ana Souza"][0]["situacao"] = "parcial"
    ex2 = r.executar("manual", ["norvia_campinas"], date(2025, 9, 1), date(2025, 9, 2))
    ana = next(g for g in r.estado.analise(ex2["empresas"]["norvia_campinas"]["analise_id"]).grupos if g.id_venda == "3000000000000001")
    assert ana.situacao == ml.SIT_CONCILIADO
    res2 = r.baixar(analise_id, [next(g.id for g in salva.grupos if g.conta_id == 10)])
    assert res2["total_ignorados"] == 1 and len(tiny.baixas) == 1


def test_resolvedor_revesta_pelo_nome():
    tiny = _Tiny()
    resolv = rot.ResolvedorTiny(tiny, lambda emp: f"tok-{emp}", Settings())
    assert resolv.ids(ml.EMPRESAS_ML["revesta"]) == (555, 900)
    fixo = rot.ResolvedorTiny(tiny, lambda emp: None, Settings(ML_REVESTA_CONTA_DESTINO_ID=1, ML_REVESTA_CATEGORIA_RECEBIMENTO_ID=2))
    assert fixo.ids(ml.EMPRESAS_ML["revesta"]) == (1, 2)


def test_resolvedor_recusa_conta_ambigua():
    class _Dup(_Tiny):
        def obter_contas_financeiras_v3(self, token):
            return [{"id": 1, "descricao": "Mercado Livre - REVESTA"}, {"id": 2, "descricao": "mercado livre - revesta"}]

    with pytest.raises(RuntimeError, match="Mais de uma conta"):
        rot.ResolvedorTiny(_Dup(), lambda emp: "tok", Settings()).ids(ml.EMPRESAS_ML["revesta"])


# ─── API ──────────────────────────────────────────────────────────────────────


@pytest.fixture
def client_rotina(tmp_path):
    from app.api import app

    r = _rotina(tmp_path)
    with patch("app.api.ml_rotina", r), patch("app.api.ml_token_tiny", return_value="tok"):
        yield TestClient(app), r


def _esperar(client, job):
    for _ in range(200):
        if job["status"] != "executando":
            return job
        time.sleep(0.05)
        job = client.get(f"/api/v1/conciliacao-ml/jobs/{job['job_id']}").json()
    raise AssertionError("tarefa não terminou")


def test_api_executar_listar_e_baixar(client_rotina):
    client, _ = client_rotina
    assert client.post("/api/v1/conciliacao-ml/executar", json={"de": "2025-09-02", "ate": "2025-09-01"}).status_code == 400
    assert client.post("/api/v1/conciliacao-ml/executar", json={"de": "2025-09-01"}).status_code == 400
    assert client.post("/api/v1/conciliacao-ml/executar", json={"empresas": ["xpto"]}).status_code == 400

    job = _esperar(client, client.post(
        "/api/v1/conciliacao-ml/executar",
        json={"empresas": ["norvia_campinas"], "de": "2025-09-01", "ate": "2025-09-02"},
    ).json())
    assert job["status"] == "concluido", job
    assert job["tipo"] == "ml_executar" and "log" not in job["resultado"]
    analise_id = job["resultado"]["empresas"]["norvia_campinas"]["analise_id"]

    execucoes = client.get("/api/v1/conciliacao-ml/execucoes").json()
    assert execucoes["total"] == 1 and "log" not in execucoes["itens"][0]
    detalhe = client.get(f"/api/v1/conciliacao-ml/execucoes/{execucoes['itens'][0]['id']}").json()
    assert len(detalhe["log"]) > 3
    assert client.get("/api/v1/conciliacao-ml/execucoes/nada").status_code == 404

    itens_analises = [i["analise_id"] for i in client.get("/api/v1/conciliacao-ml/analises?empresa=norvia_campinas").json()["itens"]]
    assert analise_id in itens_analises
    assert "consolidado_norvia_campinas" in itens_analises
    analise = client.get(f"/api/v1/conciliacao-ml/analises/{analise_id}").json()
    assert analise["resumo"]["aptos_baixa"] == 2 and len(analise["comparativo"]) > 0
    assert client.get("/api/v1/conciliacao-ml/analises/inexistente").status_code == 404
    assert client.get("/api/v1/conciliacao-ml/analises/../x").status_code == 404

    arq = client.get(f"/api/v1/conciliacao-ml/analises/{analise_id}/arquivos/vendas-recebimentos")
    assert arq.status_code == 200 and "VENDAS%20X%20RECEBIMENTOS" in arq.headers["content-disposition"]

    aptos = [g["id"] for g in analise["grupos"] if g["apto_baixa"]]
    baixa = _esperar(client, client.post(
        "/api/v1/conciliacao-ml/jobs/baixar", json={"analise_id": analise_id, "grupos": aptos, "usuario": "Fulano"}
    ).json())
    assert baixa["status"] == "concluido" and baixa["resultado"]["total_sucessos"] == 1

    st = client.get("/api/v1/conciliacao-ml/status").json()
    assert [e["id"] for e in st["empresas"]] == ["norvia_campinas", "revesta"]
    assert st["ultima_execucao"]["tipo"] == "analise" and st["hora_diaria"] == 8

    avisos = client.get("/api/v1/conciliacao-ml/avisos?depois_de=0").json()
    assert avisos["ultimo_id"] >= 1 and avisos["itens"][0]["tipo"] == "conciliacao_ml_pronta"
    assert client.get(f"/api/v1/conciliacao-ml/avisos?depois_de={avisos['ultimo_id']}").json()["itens"] == []


def test_api_executar_em_andamento_409(client_rotina):
    client, r = client_rotina
    r.em_execucao = {"execucao_id": "x"}
    assert client.post("/api/v1/conciliacao-ml/executar", json={}).status_code == 409


def test_api_baixa_sem_token_tiny_409(client_rotina):
    client, r = client_rotina
    ex = r.executar("manual", ["norvia_campinas"], date(2025, 9, 1), date(2025, 9, 1))
    with patch("app.api.ml_token_tiny", return_value=None):
        res = client.post("/api/v1/conciliacao-ml/jobs/baixar",
                          json={"analise_id": ex["empresas"]["norvia_campinas"]["analise_id"], "grupos": ["G-0"]})
    assert res.status_code == 409


# ─── Novos testes: OAuth MP, MLB, Cashback e Múltiplas Baixas Parciais ─────────


def test_oauth_mercadopago_fluxo_e_cache():
    from app.ml_apis import _MP_OAUTH_TOKEN_CACHE, obter_token_oauth_mp
    from unittest.mock import MagicMock

    _MP_OAUTH_TOKEN_CACHE.clear()
    mock_http = MagicMock()
    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {"access_token": "TEST_MP_OAUTH_TOKEN_123", "expires_in": 7200}
    mock_http.post.return_value = mock_resp

    # 1ª chamada: bate na API
    token1 = obter_token_oauth_mp("revesta", "cid_1", "csec_1", http=mock_http)
    assert token1 == "TEST_MP_OAUTH_TOKEN_123"
    assert mock_http.post.call_count == 1

    # 2ª chamada: deve vir do cache sem bater na API
    token2 = obter_token_oauth_mp("revesta", "cid_1", "csec_1", http=mock_http)
    assert token2 == "TEST_MP_OAUTH_TOKEN_123"
    assert mock_http.post.call_count == 1  # Mantém 1 chamada


def test_token_mp_prioridade_oauth_e_fallback():
    from app.ml_apis import _MP_OAUTH_TOKEN_CACHE, token_mp

    _MP_OAUTH_TOKEN_CACHE.clear()

    # Cenário 1: Com Client ID / Secret preenchidos -> busca OAuth
    cfg_oauth = Settings(
        MP_REVESTA_CLIENT_ID="id_revesta",
        MP_REVESTA_CLIENT_SECRET="sec_revesta",
        MP_REVESTA_ACCESS_TOKEN="estatico_revesta",
    )
    with patch("app.ml_apis.obter_token_oauth_mp", return_value="token_oauth_dinamico"):
        tok = token_mp("revesta", "revesta", cfg_oauth)
        assert tok == "token_oauth_dinamico"

    # Cenário 2: Sem OAuth, com token estático legado do MP
    cfg_estatico = Settings(
        MP_REVESTA_CLIENT_ID=None,
        MP_REVESTA_CLIENT_SECRET=None,
        MP_REVESTA_ACCESS_TOKEN="estatico_revesta",
    )
    tok = token_mp("revesta", "revesta", cfg_estatico)
    assert tok == "estatico_revesta"

    # Cenário 3: Sem token MP, faz fallback para token ML
    cfg_ml = Settings(
        MP_REVESTA_CLIENT_ID=None,
        MP_REVESTA_CLIENT_SECRET=None,
        MP_REVESTA_ACCESS_TOKEN=None,
        ML_REVESTA_ACCESS_TOKEN="token_do_ml_revesta",
    )
    with patch("app.ml_apis._candidatos_ml_tokens_files", return_value=[]):
        tok = token_mp("revesta", "revesta", cfg_ml)
        assert tok == "token_do_ml_revesta"


def test_normalizar_id_final_e_cashback_regex():
    assert ml.normalizar_id_final("cashback_1431390928") == ("cashback_1431390928", "1431390928")
    assert ml.normalizar_id_final("CASHBACK_987654321") == ("cashback_987654321", "987654321")
    assert ml.normalizar_id_final("2000008695371397") == ("2000008695371397", None)
    assert ml.normalizar_id_final("") == ("", None)
    assert ml.normalizar_id_final(None) == ("", None)


def test_matching_mlb_e_cruzamento_cashback():
    liberacoes_teste = [
        # Linha normal de venda (bruto 120, tarifa 20 -> líquido 100)
        _lib(SOURCE_ID="8001", NET_CREDIT_AMOUNT="100.00", GROSS_AMOUNT="120.00", MP_FEE_AMOUNT="-20.00", ORDER_ID="7001", PACK_ID="9001"),
        # Linha de cashback apontando para a operação original 8001
        {
            "DATE": "2025-09-01T10:00:00.000-03:00",
            "RECORD_TYPE": "release",
            "DESCRIPTION": "cashback",
            "SOURCE_ID": "8002",
            "EXTERNAL_REFERENCE": "8001",
            "NET_CREDIT_AMOUNT": "15.00",
            "NET_DEBIT_AMOUNT": "0.00",
            "GROSS_AMOUNT": "15.00",
            "ORDER_ID": "",
            "PACK_ID": "",
        },
    ]

    pedidos_ml_teste = [
        {
            "id": 7001,
            "pack_id": 9001,
            "status": "paid",
            "payments": [{"id": 8001}],
            "order_items": [
                {"item": {"id": "MLB3984712034", "title": "Produto Teste"}}
            ],
        }
    ]

    pedidos_tiny_teste = [
        {"ecommerce": {"numeroPedidoEcommerce": "9001"}, "cliente": {"nome": "Cliente Teste", "cpfCnpj": "123.456.789-00"}}
    ]

    movs = ml.expandir_liberacoes(liberacoes_teste)
    # 2 da venda (Recebimento + Tarifa) + 1 de cashback
    assert len(movs) == 3
    # Movimento 0: Recebimento da venda
    assert movs[0].id_final == "9001"
    assert movs[0].id_operacao_original == ""
    # Movimento 2: Cashback
    assert movs[2].id_final == "cashback_8001"
    assert movs[2].id_operacao_original == "8001"

    comparativo = ml.cruzar_vendas(movs, pedidos_ml_teste, pedidos_tiny_teste)
    # Verifica que todos herdaram o MLB e cliente através do cruzamento
    assert movs[0].mlb == "MLB3984712034"
    assert movs[0].nome_cliente == "Cliente Teste"
    assert movs[2].mlb == "MLB3984712034"
    assert movs[2].nome_cliente == "Cliente Teste"

    assert len(comparativo) == 2
    assert any(c.mlb == "MLB3984712034" for c in comparativo)


def test_multiplas_baixas_parciais_e_idempotencia():
    # Cenário: conta com saldo menor que o valor original (baixa parcial legítima)
    conta_tiny = {
        "id": 445566,
        "historico": "Pedido 9001 parcela 1/1",
        "situacao": "parcial",
        "valor": "100.00",
        "saldo": "60.00",
        "valorPago": "40.00",
    }
    # 1ª Baixa parcial: R$ 35,00
    grupo1 = ml.GrupoRecebimento(
        id="G-1",
        id_venda="9001",
        indices=[0],
        chaves=["v2|op1|2025-09-01|Recebimento|35.00|0"],
        valor_total=Decimal("35.00"),
        data="01/09/2025",
        nome_cliente="Cliente Teste",
        id_operacao="op1",
        mlb="MLB123",
    )
    ml.aplicar_conta_ao_grupo(grupo1, conta_tiny)
    ml.avaliar_conta(grupo1)
    assert grupo1.apto_baixa is True
    assert "Baixa parcial" in grupo1.mensagem

    # Se o valor passar do saldo restante (ex: R$ 70,00 quando o saldo é R$ 60,00) -> divergente
    grupo_excesso = ml.GrupoRecebimento(
        id="G-2",
        id_venda="9001",
        indices=[1],
        chaves=["v2|op2|2025-09-01|Recebimento|70.00|0"],
        valor_total=Decimal("70.00"),
        data="01/09/2025",
        nome_cliente="Cliente Teste",
        id_operacao="op2",
        mlb="MLB123",
    )
    ml.aplicar_conta_ao_grupo(grupo_excesso, conta_tiny)
    ml.avaliar_conta(grupo_excesso)
    assert grupo_excesso.situacao == ml.SIT_VALOR_DIVERGENTE


def test_gerar_arquivo_relatorio_colunas_concatenadas():
    analise = _analisar()
    nome, bytes_arq = ml.gerar_arquivo(analise, "relatorio")
    assert nome.endswith(".xlsx")

    wb = load_workbook(io.BytesIO(bytes_arq), read_only=True)
    ws = wb.active
    cabecalho = list(next(ws.iter_rows(values_only=True)))

    # As 8 primeiras colunas devem bater rigorosamente com a planilha RELATÓRIO CONCATENADO AGOSTO REVESTA 2025.xlsx:
    assert cabecalho[:8] == [
        "Data de pagamento",
        "Tipo de operação",
        "Número do movimento",
        "Operação relacionada",
        "ID_FINAL",
        "Nome do cliente",
        "Valor",
        "Situação",
    ]
    # 9ª coluna deve ser MLB
    assert cabecalho[8] == "MLB"
    wb.close()


def test_venda_com_frete_unificada_e_baixa_categorias_distintas():
    # 1. Relatório com 1 linha de produto e 1 linha de envio do mesmo pedido/pacote
    liberacoes = [
        {
            "RECORD_TYPE": "release",
            "DESCRIPTION": "shipping",
            "DATE": "10/08/2026",
            "SOURCE_ID": "168827487964",
            "ORDER_ID": "2000017418947554",
            "PACK_ID": "2000014021072793",
            "NET_CREDIT_AMOUNT": "44.05",
            "NET_DEBIT_AMOUNT": "0.00",
            "GROSS_AMOUNT": "44.05",
        },
        {
            "RECORD_TYPE": "release",
            "DESCRIPTION": "payment",
            "DATE": "10/08/2026",
            "SOURCE_ID": "168827773708",
            "ORDER_ID": "2000017418947554",
            "PACK_ID": "2000014021072793",
            "NET_CREDIT_AMOUNT": "133.78",
            "NET_DEBIT_AMOUNT": "0.00",
            "GROSS_AMOUNT": "133.78",
        },
    ]
    movs = ml.expandir_liberacoes(liberacoes)
    assert len(movs) == 2
    assert movs[0].tipo_operacao == ml.TIPO_RECEBIMENTO_FRETE
    assert movs[0].valor == Decimal("44.05")
    assert movs[1].tipo_operacao == ml.TIPO_RECEBIMENTO
    assert movs[1].valor == Decimal("133.78")

    # Pedido no ML com o pack_id
    pedidos_ml = [
        {
            "id": 2000017418947554,
            "pack_id": 2000014021072793,
            "status": "paid",
            "payments": [{"id": 168827773708}],
            "order_items": [{"item": {"id": "MLB5211415434"}}],
        }
    ]
    pedidos_tiny = [
        {
            "ecommerce": {"numeroPedidoEcommerce": "2000014021072793"},
            "cliente": {"nome": "Alfredo Exemplo", "cpfCnpj": "123.456.789-00"},
        }
    ]

    comp = ml.cruzar_vendas(movs, pedidos_ml, pedidos_tiny)
    # Deve unificar em 1 linha no comparativo da venda!
    assert len(comp) == 1
    assert comp[0].valor_liberado == Decimal("177.83")
    assert comp[0].nome_cliente == "Alfredo Exemplo"
    assert comp[0].id_pacote == "2000014021072793"

    # Lançamentos de caixa: nenhum lançamento de caixa avulso deve ser gerado para o frete!
    resultado_lanc = ml.gerar_lancamentos(movs, "Mercado Livre - NORVIA")
    assert len(resultado_lanc.lancamentos) == 0

    # Grupo de recebimento: deve somar produto + frete
    grupos = ml.obter_grupos(movs)
    assert len(grupos) == 1
    g = grupos[0]
    assert g.valor_total == Decimal("177.83")
    assert g.valor_produtos == Decimal("133.78")
    assert g.valor_frete == Decimal("44.05")

    # Conta no Tiny de R$ 177,83
    conta_tiny = {
        "id": 895293685,
        "valor": 177.83,
        "saldo": 177.83,
        "valorPago": 0.0,
        "situacao": "aberto",
        "historico": "Ref. ao pedido de venda nº 32171, Alfredo Exemplo - OC nº 2000014021072793",
    }
    ml.aplicar_conta_ao_grupo(g, conta_tiny)
    ml.avaliar_conta(g)

    assert g.situacao == ml.SIT_APTO_BAIXA
    assert g.mensagem == ""  # Não deve dizer 'Baixa parcial: ficam R$ 44,05 em aberto'!


def test_rotina_baixar_com_frete_chama_categorias_distintas():
    from app.ml_reconciliation_rotina import RotinaConciliacaoML, ResolvedorTiny

    empresa = ml.EMPRESAS_ML["norvia_campinas"]
    tiny_mock = MagicMock()
    tiny_mock.obter_conta_receber_v3.return_value = {
        "id": 895293685,
        "valor": 177.83,
        "saldo": 177.83,
        "valorPago": 0.0,
        "situacao": "aberto",
    }
    estado_mock = MagicMock()
    estado_mock.ler.return_value = {}

    rotina = RotinaConciliacaoML(estado=estado_mock, tiny=tiny_mock, cfg=Settings())

    # Mock da análise
    grupo = ml.GrupoRecebimento(
        id="G-1",
        id_venda="2000014021072793",
        indices=[0, 1],
        chaves=["k1", "k2"],
        valor_total=Decimal("177.83"),
        valor_produtos=Decimal("133.78"),
        valor_frete=Decimal("44.05"),
        data="10/08/2026",
        nome_cliente="Alfredo Exemplo",
        conta_id=895293685,
        conta_situacao="aberto",
        conta_saldo=Decimal("177.83"),
        situacao=ml.SIT_APTO_BAIXA,
    )
    analise = ml.Analise(
        id="analise_teste",
        empresa_id="norvia_campinas",
        empresa_nome="Norvia Campinas",
        conta_lancamento="Mercado Livre - NORVIA",
        criada_em="2026-10-09T00:00:00",
        periodo_de="2026-08-10",
        periodo_ate="2026-08-10",
        origem="manual",
        execucao_id="ex1",
        movimentos=[],
        grupos=[grupo],
        lancamentos=ml.ResultadoLancamentos(),
        comparativo=[],
        fontes={},
        tiny_consultado=True,
    )
    estado_mock.analise.return_value = analise
    rotina.obter_token_tiny = MagicMock(return_value="fake_token")

    res = rotina.baixar("analise_teste", ["G-1"])
    assert res["status"] == "concluido"
    assert res["total_sucessos"] == 1

    # Verifica as duas chamadas a baixar_conta_receber_v3:
    # 1. Parcela de produtos com categoria de vendas (882461346)
    # 2. Parcela de frete com categoria de frete (489820971)
    chamadas = tiny_mock.baixar_conta_receber_v3.call_args_list
    assert len(chamadas) == 2

    # Primeira chamada: produtos
    args1, kwargs1 = chamadas[0]
    assert args1[1] == 895293685
    assert kwargs1["valor_pago"] == 133.78
    assert kwargs1["categoria_id"] == 882461346

    # Segunda chamada: frete
    args2, kwargs2 = chamadas[1]
    assert args2[1] == 895293685
    assert kwargs2["valor_pago"] == 44.05
    assert kwargs2["categoria_id"] == 489820971


def test_pack_id_indexado_e_fallback_ml_api():
    from app.ml_apis import ClienteMercadoLivre, APIMercadoError

    pedidos = [{"id": 111, "pack_id": 999, "payments": []}]
    por_pag, por_ped = ml.indice_pedidos_ml(pedidos)
    assert "111" in por_ped
    assert "999" in por_ped

    # Fallback no ClienteMercadoLivre quando /orders/{id} dá 404
    cliente = ClienteMercadoLivre("fake_conta", lambda: "fake_token")
    def mock_chamar(metodo, path, descricao, **kwargs):
        if path == "/orders/999":
            raise APIMercadoError("Not Found", status=404)
        if path == "/packs/999":
            return {"id": 999, "orders": [{"id": 111}]}
        if path == "/orders/111":
            return {"id": 111, "status": "paid"}
        raise APIMercadoError("Not Found", status=404)

    cliente._chamar = mock_chamar
    ped = cliente.pedido("999")
    assert ped is not None
    assert ped.get("id") == 111
    assert ped.get("pack_id") == "999"




# ─── custos de parcelamento (venda parcelada com juros do comprador) ─────────────


def _venda_parcelada(saldo_tiny):
    """Caso real (24/08/2026): recebimento 116,94 + frete 69,90 = 186,84; tarifa -10,33; custos de
    parcelamento -21,14; a conta do Tiny tem o preço da venda."""
    libs = [
        _lib(SOURCE_ID="7001", NET_CREDIT_AMOUNT="155.37", GROSS_AMOUNT="186.84", MP_FEE_AMOUNT="-10.33",
             FINANCING_FEE_AMOUNT="-21.14", ORDER_ID="2000000000000001"),
    ]
    pedidos_ml = [{"id": 2000000000000001, "pack_id": 3000000000000001, "status": "paid", "payments": [{"id": 7001}]}]
    conta = {"id": 77, "situacao": "aberto", "valor": saldo_tiny, "saldo": saldo_tiny,
             "historico": "Ref. ao pedido de venda - OC nº 3000000000000001"}
    return ml.analisar("p1", NORVIA, libs, pedidos_ml, PEDIDOS_TINY, lambda nome: [conta],
                       (date(2025, 9, 1), date(2025, 9, 2)))


def test_grupo_soma_custos_de_parcelamento_da_venda():
    analise = _venda_parcelada("165.70")
    (g,) = [g for g in analise.grupos if g.id_venda == "3000000000000001"]
    assert g.valor_parcelamento == Decimal("21.14") and len(g.indices_parcelamento) == 1


def test_parcelado_com_juros_vira_apto_pelo_valor_do_tiny():
    analise = _venda_parcelada("165.70")
    (g,) = [g for g in analise.grupos if g.id_venda == "3000000000000001"]
    assert g.apto_baixa and g.parcelamento_descontado
    assert g.valor_total == Decimal("165.70") and g.valor_produtos + g.valor_frete == Decimal("165.70")
    assert "custos de parcelamento" in g.mensagem
    # O custo descontado na baixa não vira lançamento de caixa (a saída entraria duas vezes);
    # a tarifa de venda continua lançada.
    tipos = [l.tipo_operacao for l in analise.lancamentos.lancamentos]
    assert "Custos de parcelamento" not in tipos and "Tarifa de venda" in tipos
    assert not [p for p in analise.lancamentos.pendencias if p.tipo_operacao == "Custos de parcelamento"]
    # Avaliar de novo (visão consolidada, baixa) não desconta duas vezes.
    ml.avaliar_conta(g)
    assert g.valor_total == Decimal("165.70") and g.apto_baixa


def test_parcelado_sem_juros_tiny_com_bruto_mantem_custo_como_despesa():
    analise = _venda_parcelada("186.84")
    (g,) = [g for g in analise.grupos if g.id_venda == "3000000000000001"]
    assert g.apto_baixa and not g.parcelamento_descontado and g.valor_total == Decimal("186.84")
    assert "Custos de parcelamento" in [l.tipo_operacao for l in analise.lancamentos.lancamentos]


def test_diferenca_que_nao_e_o_parcelamento_continua_divergente():
    analise = _venda_parcelada("150.00")
    (g,) = [g for g in analise.grupos if g.id_venda == "3000000000000001"]
    assert g.situacao == ml.SIT_VALOR_DIVERGENTE and not g.parcelamento_descontado
    assert g.valor_total == Decimal("186.84")


def test_nome_cliente_longo_vai_cortado_em_50_para_o_tiny():
    # Tiny v3 responde 400 ("Deve ter 50 caracteres ou menos") com nomeCliente maior.
    from app.tiny_client import TinyClient

    client = TinyClient(token="mock_token")
    nome = "Instituto Vital Brazil Sa (centro de Pesquisas, Produtos Quimicos e Biologicos)"
    with patch.object(TinyClient, "_fetch_paginated_v3_itens", return_value=[]) as fetch:
        client.list_contas_receber_v3("tok", nome_cliente=nome, estrito=True)
    enviado = fetch.call_args.args[2]["nomeCliente"]
    assert len(enviado) <= 50 and nome.startswith(enviado)
