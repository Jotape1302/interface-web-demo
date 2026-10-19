"""Correções da revisão de 2026-10-07 na conciliação Mercado Livre (ML-1 a ML-7)."""
import io
from datetime import date
from decimal import Decimal
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import ml_reconciliation as ml
from app import ml_reconciliation_rotina as rot
from tests.test_ml_reconciliation import LIBERACOES, _Tiny, _lib, _rotina

DIA = (date(2025, 9, 1), date(2025, 9, 2))


def _executar(r):
    ex = r.executar("manual", ["norvia_campinas"], *DIA)
    analise = r.estado.analise(ex["empresas"]["norvia_campinas"]["analise_id"])
    return analise, [g.id for g in analise.grupos if g.apto_baixa]


# ─── ML-1: chave estável entre reprocessamentos ──────────────────────────────


def test_ml1_chave_nao_depende_da_linha_do_relatorio():
    venda = _lib(SOURCE_ID="7001", NET_CREDIT_AMOUNT="90.00", GROSS_AMOUNT="100.00", MP_FEE_AMOUNT="-10.00")
    outra = _lib(SOURCE_ID="7002", NET_CREDIT_AMOUNT="50.00", GROSS_AMOUNT="50.00")
    a = {m.tipo_operacao: m.chave for m in ml.expandir_liberacoes([venda, outra]) if m.id_pagamento == "7001"}
    # Mesmo período reprocessado com outras datas: a liberação cai em outra linha do relatório.
    b = {m.tipo_operacao: m.chave for m in ml.expandir_liberacoes([outra, outra | {"SOURCE_ID": "7003"}, venda]) if m.id_pagamento == "7001"}
    assert a == b and a[ml.TIPO_RECEBIMENTO].startswith("v2|7001|")


def test_ml1_baixa_registrada_com_chave_antiga_continua_valendo():
    movs = ml.expandir_liberacoes([_lib(SOURCE_ID="7001", NET_CREDIT_AMOUNT="100.00", GROSS_AMOUNT="100.00", ORDER_ID="8001")])
    rec = next(m for m in movs if m.tipo_operacao == ml.TIPO_RECEBIMENTO)
    chave_antiga = f"7001|{rec.data}|{rec.tipo_operacao}|100.00|17"  # formato antigo, linha 17 de outro relatório
    grupos = ml.obter_grupos(movs, {chave_antiga: "2026-10-01"})
    assert grupos[0].situacao == ml.SIT_CONCILIADO and grupos[0].baixado_em == "2026-10-01"


def test_ml1_reprocessar_periodo_nao_baixa_de_novo_conta_parcial(tmp_path):
    class TinyParcial(_Tiny):
        """Conta de R$ 220 (duas liberações de R$ 110): depois da 1ª baixa fica parcial."""

        def __init__(self):
            super().__init__()
            self.contas["Ana Souza"][0].update(valor=220, saldo=220)
            self.saldo = {10: 220}

        def obter_conta_receber_v3(self, token, cid):
            if cid != 10:
                return super().obter_conta_receber_v3(token, cid)
            s = self.saldo[10]
            return {"situacao": "parcial" if s < 220 else "aberto", "valor": 220, "saldo": s, "valorPago": 220 - s}

        def baixar_conta_receber_v3(self, token, cid, **kw):
            self.baixas.append((cid, kw))
            if cid == 10:
                self.saldo[10] -= kw["valor_pago"]

    tiny = TinyParcial()
    r = _rotina(tmp_path, tiny=tiny)
    analise, aptos = _executar(r)
    grupo_ana = next(g.id for g in analise.grupos if g.conta_id == 10)
    r.baixar(analise.id, [grupo_ana])
    assert [c for c, _ in tiny.baixas] == [10]

    # Reprocessa com o relatório em outra ordem (a liberação muda de linha).
    r2 = _rotina(tmp_path, tiny=tiny)
    r2.fabrica_mp = lambda e: type("MP", (), {"liberacoes": lambda self, de, ate, progresso=None: LIBERACOES[::-1]})()
    analise2, _ = _executar(r2)
    g2 = next(g for g in analise2.grupos if g.id_venda == next(g.id_venda for g in analise.grupos if g.conta_id == 10))
    assert g2.situacao == ml.SIT_CONCILIADO
    r2.baixar(analise2.id, [g2.id])
    assert [c for c, _ in tiny.baixas] == [10]  # nenhuma baixa nova


# ─── ML-2: resposta incerta do Tiny ─────────────────────────────────────────


class _TinyComValorPago(_Tiny):
    def __init__(self, falha=None, entra=True, releitura_falha=False):
        super().__init__()
        self.pago = {10: 0.0, 20: 80.0}
        self.falha, self.entra, self.releitura_falha = falha, entra, releitura_falha
        self.leituras = 0

    def obter_conta_receber_v3(self, token, cid):
        self.leituras += 1
        if self.releitura_falha and self.baixas:
            raise ConnectionError("Tiny fora do ar")
        valor = 110 if cid == 10 else 80
        pago = self.pago[cid]
        return {"situacao": "pago" if pago >= valor else "aberto", "valor": valor, "saldo": valor - pago, "valorPago": pago}

    def baixar_conta_receber_v3(self, token, cid, **kw):
        self.baixas.append((cid, kw))
        if self.entra:
            self.pago[cid] += kw["valor_pago"]
        if self.falha:
            raise TimeoutError(self.falha)


def _grupo_ana(analise):
    return next(g.id for g in analise.grupos if g.conta_id == 10)


def test_ml2_timeout_mas_baixa_entrou_conta_como_sucesso_e_registra(tmp_path):
    tiny = _TinyComValorPago(falha="read timeout", entra=True)
    r = _rotina(tmp_path, tiny=tiny)
    analise, _ = _executar(r)
    res = r.baixar(analise.id, [_grupo_ana(analise)])
    assert res["total_sucessos"] == 1 and res["total_falhas"] == 0
    dados = r.estado.ler()
    assert dados["baixas_pendentes"] == {} and len(dados["baixas"]) >= 1


def test_ml2_erro_sem_baixa_libera_nova_tentativa(tmp_path):
    tiny = _TinyComValorPago(falha="HTTP 500", entra=False)
    r = _rotina(tmp_path, tiny=tiny)
    analise, _ = _executar(r)
    res = r.baixar(analise.id, [_grupo_ana(analise)])
    assert res["total_falhas"] == 1 and r.estado.ler()["baixas_pendentes"] == {}
    tiny.falha, tiny.entra = None, True
    assert r.baixar(analise.id, [_grupo_ana(analise)])["total_sucessos"] == 1
    assert len(tiny.baixas) == 2 and tiny.pago[10] == 110


def test_ml2_sem_confirmacao_a_proxima_tentativa_confere_antes_de_baixar(tmp_path):
    tiny = _TinyComValorPago(falha="read timeout", entra=True, releitura_falha=True)
    r = _rotina(tmp_path, tiny=tiny)
    analise, _ = _executar(r)
    res = r.baixar(analise.id, [_grupo_ana(analise)])
    assert res["total_falhas"] == 1 and "sem confirmação" in res["resultados"][0]["mensagem"]
    assert len(r.estado.ler()["baixas_pendentes"]) == 1

    # Tiny volta: a tentativa anterior tinha entrado — não baixa de novo.
    tiny.releitura_falha, tiny.falha = False, None
    res = r.baixar(analise.id, [_grupo_ana(analise)])
    assert res["total_ignorados"] == 1 and len(tiny.baixas) == 1 and tiny.pago[10] == 110
    assert r.estado.ler()["baixas_pendentes"] == {}


def test_ml2_disco_falhando_antes_da_baixa_nao_chama_o_tiny(tmp_path):
    tiny = _TinyComValorPago()
    r = _rotina(tmp_path, tiny=tiny)
    analise, _ = _executar(r)
    with patch.object(r.estado, "alterar", side_effect=OSError("disco cheio")):
        with pytest.raises(OSError):
            r.baixar(analise.id, [_grupo_ana(analise)])
    assert tiny.baixas == []


# ─── ML-3: trava por empresa ────────────────────────────────────────────────


def test_ml3_baixas_simultaneas_da_mesma_empresa_sao_recusadas(tmp_path):
    r = _rotina(tmp_path)
    analise, aptos = _executar(r)
    r._baixas_em_andamento.add("norvia_campinas")  # outra análise baixando a mesma empresa
    with pytest.raises(rot.EmExecucao):
        r.baixar(analise.id, aptos)


# ─── ML-4 / ML-7: nada some em silêncio ─────────────────────────────────────


def test_ml4_valor_ilegivel_vira_pendencia_e_nao_entra_em_anulacao():
    regs = [
        _lib(SOURCE_ID="7100", DESCRIPTION="refund", NET_DEBIT_AMOUNT="10.00"),
        _lib(SOURCE_ID="7100", DESCRIPTION="refund_cancel", NET_CREDIT_AMOUNT="10.00"),
        _lib(SOURCE_ID="7100", DESCRIPTION="shipping", NET_DEBIT_AMOUNT="1O,50"),  # letra O no lugar de zero
    ]
    movs = ml.expandir_liberacoes(regs)
    ilegivel = next(m for m in movs if m.tipo_operacao == ml.TIPO_VALOR_ILEGIVEL)
    assert "1O,50" in ilegivel.descricao_mp
    res = ml.gerar_lancamentos(movs, "Conta")
    assert ilegivel.situacao == ml.SIT_VERIFICAR
    assert any("ilegível" in p.motivo for p in res.pendencias)
    with pytest.raises(ValueError):
        ml.valor_estrito("abc")
    assert ml.para_decimal("abc") == 0  # campos informativos continuam tolerantes


def test_ml7_linha_repetida_nao_some_nem_vira_lancamento():
    tarifa = _lib(SOURCE_ID="7200", DESCRIPTION="shipping", NET_DEBIT_AMOUNT="5.00")
    movs = ml.expandir_liberacoes([tarifa, dict(tarifa)])
    assert [m.tipo_operacao for m in movs] == ["Tarifa de envio", ml.TIPO_LINHA_REPETIDA]
    res = ml.gerar_lancamentos(movs, "Conta")
    assert len(res.lancamentos) == 1 and any("idêntica" in p.motivo for p in res.pendencias)


# ─── ML-6: saldo desconhecido ───────────────────────────────────────────────


def test_ml6_conta_sem_saldo_nem_valor_nao_fica_apta():
    g = ml.GrupoRecebimento(id="G-1", id_venda="1", indices=[0], chaves=["k"], valor_total=Decimal("50"),
                            data="01/09/2025", nome_cliente="X")
    ml.aplicar_conta_ao_grupo(g, {"id": 1, "situacao": "aberto"})
    ml.avaliar_conta(g)
    assert g.situacao == ml.SIT_SALDO_DESCONHECIDO and not g.apto_baixa


# ─── ML-5: lançamentos já exportados ────────────────────────────────────────


def test_ml5_planilha_de_lancamentos_nao_repete_o_que_ja_foi_exportado(tmp_path):
    from app.api import app

    r = _rotina(tmp_path)
    with patch("app.api.ml_rotina", r), patch("app.api.ml_token_tiny", return_value="tok"):
        client = TestClient(app)
        analise, _ = _executar(r)
        url = f"/api/v1/conciliacao-ml/analises/{analise.id}/arquivos/lancamentos"
        total = len(analise.lancamentos.lancamentos)
        assert total > 0

        primeira = client.get(url)
        assert primeira.status_code == 200
        assert primeira.headers["x-lancamentos-na-planilha"] == str(total)
        assert load_workbook(io.BytesIO(primeira.content)).active.max_row == total + 1

        segunda = client.get(url)
        assert segunda.status_code == 409 and "já foram exportados" in segunda.json()["detail"]

        todos = client.get(url + "?incluir_exportados=true")
        assert todos.status_code == 200 and todos.headers["x-lancamentos-ja-exportados"] == str(total)
        assert load_workbook(io.BytesIO(todos.content)).active.max_row == total + 1


# ─── Matching de reservas, mensagens claras e transferências ────────────────


def test_anulacao_reserva_em_venda_com_tarifas():
    # Venda que possui tarifas normais E retenção/liberação juntas no mesmo SOURCE_ID
    regs = [
        _lib(SOURCE_ID="8001", DESCRIPTION="payment", GROSS_AMOUNT="100.00", NET_CREDIT_AMOUNT="80.00",
             MP_FEE_AMOUNT="-12.00", SHIPPING_FEE_AMOUNT="-8.00", ORDER_ID="2000000000000001"),
        _lib(SOURCE_ID="8001", DESCRIPTION="reserve_for_shipping", NET_DEBIT_AMOUNT="14.24"),
        _lib(SOURCE_ID="8001", DESCRIPTION="release_reserve_shipping", NET_CREDIT_AMOUNT="14.24"),
    ]
    movs = ml.expandir_liberacoes(regs)
    # Associa venda e cliente para não falhar por falta de cliente nas tarifas
    for m in movs:
        if m.id_pagamento == "8001":
            m.nome_cliente = "Ana Souza"
            m.cpf_cnpj = "111.111.111-11"

    res = ml.gerar_lancamentos(movs, "Conta ML")
    assert res.total_anuladas == 2
    # Nenhuma pendência para as reservas
    assert not any("14.24" in str(p.valor) for p in res.pendencias)
    # As tarifas normais (12.00 e 8.00) viraram lançamentos
    valores_lanc = [l.valor for l in res.lancamentos]
    assert Decimal("12.00") in valores_lanc
    assert Decimal("8.00") in valores_lanc


def test_reserva_sem_par_com_motivo_claro():
    # Retenção isolada
    reg_deb = _lib(SOURCE_ID="8002", DESCRIPTION="reserve_for_payment", NET_DEBIT_AMOUNT="25.50")
    movs_deb = ml.expandir_liberacoes([reg_deb])
    res_deb = ml.gerar_lancamentos(movs_deb, "Conta")
    assert len(res_deb.pendencias) == 1
    assert "aguardando liberação futura" in res_deb.pendencias[0].motivo
    assert res_deb.pendencias[0].fluxo == "PAGAMENTO"

    # Liberação isolada
    reg_crd = _lib(SOURCE_ID="8003", DESCRIPTION="release_reserve_payment", NET_CREDIT_AMOUNT="25.50")
    movs_crd = ml.expandir_liberacoes([reg_crd])
    res_crd = ml.gerar_lancamentos(movs_crd, "Conta")
    assert len(res_crd.pendencias) == 1
    assert "retenção ocorreu em período anterior" in res_crd.pendencias[0].motivo
    assert res_crd.pendencias[0].fluxo == "RECEBIMENTO"


def test_transferencias_e_pix_entre_contas_grupo():
    for desc in ("transfer", "money_transfer", "bank_transfer", "internal_transfer", "pix", "pix_transfer"):
        assert ml.tipo_por_descricao(desc, False) == ml.TIPO_TRANSFERENCIA
        assert ml.tipo_por_descricao(desc, True) == ml.TIPO_TRANSFERENCIA


def test_motivos_especificos_sem_par_tiny_ou_ml():
    # Movimento sem venda ligada
    mov_sem_venda = ml.Movimento(
        indice=0, data="01/09/2025", tipo_operacao="Tarifa de venda", descricao_mp="fee",
        id_pagamento="9001", id_pedido="", id_pacote="", valor=Decimal("-10.00"), linha_relatorio=1
    )
    res1 = ml.gerar_lancamentos([mov_sem_venda], "Conta")
    assert "Não encontrado par no Mercado Livre" in res1.pendencias[0].motivo

    # Movimento com venda no ML mas não achada no Tiny
    mov_sem_tiny = ml.Movimento(
        indice=0, data="01/09/2025", tipo_operacao="Tarifa de venda", descricao_mp="fee",
        id_pagamento="9002", id_pedido="20001", id_pacote="", valor=Decimal("-10.00"), linha_relatorio=1
    )
    res2 = ml.gerar_lancamentos([mov_sem_tiny], "Conta")
    assert "Não encontrado par no Tiny" in res2.pendencias[0].motivo


# ─── Vendas parceladas: agrupamento e baixa consolidada ──────────────────────


def test_venda_parcelada_2x_agrupa_parcelas_e_fica_apta():
    venda_id = "2000014584857665"
    contas = [
        {
            "id": 824398571,
            "situacao": "aberto",
            "valor": 31.47,
            "saldo": 31.47,
            "valorPago": 0,
            "historico": f"Venda {venda_id} parcela 1/2",
        },
        {
            "id": 824398572,
            "situacao": "aberto",
            "valor": 31.47,
            "saldo": 31.47,
            "valorPago": 0,
            "historico": f"Venda {venda_id} parcela 2/2",
        },
    ]

    achadas = ml.selecionar_contas_da_venda(contas, venda_id)
    assert len(achadas) == 2
    assert achadas[0]["id"] == 824398571
    assert achadas[1]["id"] == 824398572

    g = ml.GrupoRecebimento(
        id="G-1",
        id_venda=venda_id,
        indices=[0],
        chaves=["k1"],
        valor_total=Decimal("62.95"),
        data="08/10/2026",
        nome_cliente="Joana Mendes",
    )
    ml.aplicar_contas_ao_grupo(g, achadas)

    assert len(g.contas_vinculadas) == 2
    assert g.conta_saldo == Decimal("62.94")
    assert g.conta_valor == Decimal("62.94")
    assert "2 parcelas" in g.conta_historico

    ml.avaliar_conta(g)
    assert g.situacao == ml.SIT_APTO_BAIXA
    assert g.apto_baixa
    assert "2 parcelas" in g.mensagem

    # Serialização JSON inclui as parcelas
    json_data = ml.grupo_json(g)
    assert json_data["quantidade_parcelas"] == 2
    assert len(json_data["contas_vinculadas"]) == 2
    assert json_data["contas_vinculadas"][0]["id"] == 824398571


def test_venda_parcelada_3x_agrupa_parcelas():
    venda_id = "2000014684449453"
    contas = [
        {"id": 824407117, "situacao": "aberto", "valor": 20.14, "saldo": 20.14, "historico": f"Venda {venda_id} parcela 1/3"},
        {"id": 824407118, "situacao": "aberto", "valor": 20.14, "saldo": 20.14, "historico": f"Venda {venda_id} parcela 2/3"},
        {"id": 824407119, "situacao": "aberto", "valor": 20.14, "saldo": 20.14, "historico": f"Venda {venda_id} parcela 3/3"},
    ]

    g = ml.GrupoRecebimento(
        id="G-2",
        id_venda=venda_id,
        indices=[0],
        chaves=["k2"],
        valor_total=Decimal("60.42"),
        data="08/10/2026",
        nome_cliente="Mateus Ramos",
    )
    ml.aplicar_contas_ao_grupo(g, contas)
    ml.avaliar_conta(g)

    assert g.situacao == ml.SIT_APTO_BAIXA
    assert g.apto_baixa
    assert g.conta_saldo == Decimal("60.42")
    assert "3 parcelas" in g.mensagem


def test_rotina_baixa_distribui_valor_entre_todas_as_parcelas(tmp_path):
    from unittest.mock import MagicMock
    from app.ml_reconciliation_rotina import RotinaConciliacaoML
    from app.config import Settings
    from app.ml_reconciliation_rotina import EstadoML

    cfg = Settings(APP_DATA_DIR=str(tmp_path), DATA_DIR=str(tmp_path / "data"))
    estado = EstadoML(pasta=tmp_path / "conciliacao_ml", cfg=cfg)
    rotina = RotinaConciliacaoML(estado=estado, cfg=cfg, obter_token_tiny=lambda _: "tok_teste")

    # Mock do TinyClient
    contas_db = {
        824398571: {"id": 824398571, "situacao": "aberto", "valor": 31.47, "saldo": 31.47, "valorPago": 0, "historico": "Venda 2000014584857665 parcela 1/2"},
        824398572: {"id": 824398572, "situacao": "aberto", "valor": 31.47, "saldo": 31.47, "valorPago": 0, "historico": "Venda 2000014584857665 parcela 2/2"},
    }
    baixas_chamadas = []

    def fake_obter(token, cid):
        return dict(contas_db[cid])

    def fake_baixar(token, cid, valor_pago, data_pagamento, conta_destino_id, categoria_id):
        baixas_chamadas.append({"cid": cid, "valor_pago": valor_pago})
        contas_db[cid]["saldo"] = max(0.0, contas_db[cid]["saldo"] - valor_pago)
        contas_db[cid]["valorPago"] += valor_pago
        if contas_db[cid]["saldo"] == 0:
            contas_db[cid]["situacao"] = "pago"

    rotina._tiny = MagicMock()
    rotina._tiny.obter_conta_receber_v3.side_effect = fake_obter
    rotina._tiny.baixar_conta_receber_v3.side_effect = fake_baixar

    # Cria análise com grupo parcelado
    empresa = ml.EMPRESAS_ML["revesta"]
    g = ml.GrupoRecebimento(
        id="G-100",
        id_venda="2000014584857665",
        indices=[0],
        chaves=["chave_mov_1"],
        valor_total=Decimal("62.94"),
        data="08/10/2026",
        nome_cliente="Joana",
    )
    ml.aplicar_contas_ao_grupo(g, list(contas_db.values()))
    ml.avaliar_conta(g)
    assert g.apto_baixa

    mov = ml.Movimento(
        indice=0,
        tipo_operacao="Recebimento de venda",
        valor=Decimal("62.94"),
        data="08/10/2026",
        descricao_mp="payment",
        id_pedido="2000014584857665",
        id_pacote="",
        id_pagamento="999",
        linha_relatorio=1,
    )
    analise = ml.Analise(
        id="analise-teste",
        empresa_id=empresa.id,
        empresa_nome=empresa.nome,
        conta_lancamento=empresa.conta_lancamento,
        criada_em="2026-10-09T00:00:00",
        periodo_de="2026-10-08",
        periodo_ate="2026-10-08",
        origem="manual",
        execucao_id="exec-1",
        movimentos=[mov],
        grupos=[g],
        lancamentos=ml.ResultadoLancamentos([], [], 0),
        comparativo=[],
        fontes={},
        tiny_consultado=True,
    )
    rotina.estado.salvar_analise(analise)

    # Resolvedor mock
    rotina._resolvedor = MagicMock()
    rotina._resolvedor.ids.return_value = (100, 200)

    # Executa a baixa
    resultado = rotina.baixar("analise-teste", [g.id], usuario="admin")
    assert resultado["status"] == "concluido"
    assert resultado["total_sucessos"] == 1
    assert len(baixas_chamadas) == 2
    assert baixas_chamadas[0] == {"cid": 824398571, "valor_pago": 31.47}
    assert baixas_chamadas[1] == {"cid": 824398572, "valor_pago": 31.47}

    # Verifica situação do grupo após baixa
    analise_depois = rotina.estado.analise("analise-teste")
    assert analise_depois.grupos[0].situacao == ml.SIT_CONCILIADO
    assert "2 parcelas" in resultado["resultados"][0]["mensagem"]


